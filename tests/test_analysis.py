import pandas as pd

from flutracker import analysis, emailer


def rows(country, file, values, **extra):
    return [
        {"file": file, "countryname": country, "yearweek": yw, "value": v, **extra}
        for yw, v in values.items()
    ]


def season_values(start, base):
    """A simple season: `base` at every week from W40 to W39."""
    return {analysis.yw_str(y, w): base for y, w in analysis.season_weeks(start)}


def make_raw():
    lab = dict(pathogen="Influenza", pathogentype="Influenza", pathogensubtype="total", age="total", survtype="non-sentinel")
    covid = dict(pathogen="SARS-CoV-2", pathogentype="SARS-CoV-2", pathogensubtype="SARS-CoV-2", age="total", survtype="non-sentinel")
    data = []
    for start, base in [(2022, 10), (2023, 20), (2024, 30)]:
        data += rows("Norway", "nonSentinelTestsDetections", season_values(start, base), indicator="detections", **lab)
        data += rows("Norway", "nonSentinelTestsDetections", season_values(start, 100), indicator="tests", **lab)
        data += rows("Norway", "nonSentinelTestsDetections", season_values(start, base * 10), indicator="detections", **covid)
        data += rows("Sweden", "nonSentinelSeverity", season_values(start, base), pathogen="Influenza",
                     pathogentype="Influenza", indicator="hospitaladmissions", age="total", survtype="non-sentinel")
    current = {"2025-W40": 40, "2025-W41": 50}
    data += rows("Norway", "nonSentinelTestsDetections", current, indicator="detections", **lab)
    data += rows("Norway", "nonSentinelTestsDetections", {k: 100 for k in current}, indicator="tests", **lab)
    data += rows("Norway", "nonSentinelTestsDetections", {k: 100 for k in current}, indicator="detections", **covid)
    data += rows("Sweden", "nonSentinelSeverity", {"2025-W41": 15}, pathogen="Influenza",
                 pathogentype="Influenza", indicator="hospitaladmissions", age="total", survtype="non-sentinel")
    data += rows("Norway", "activityFluTypeSubtype", {"2025-W41": 7}, pathogen="Influenza",
                 pathogentype="Influenza A", pathogensubtype="A(H3)", indicator="detections", age="total", survtype="non-sentinel")
    return pd.DataFrame(data)


def test_season_helpers():
    assert analysis.season_start(2026, 39) == 2025
    assert analysis.season_start(2026, 40) == 2026
    assert analysis.season_label(2026) == "2026/27"
    weeks = analysis.season_weeks(2026)  # 2026 has ISO week 53
    assert weeks[0] == (2026, 40) and (2026, 53) in weeks and weeks[-1] == (2027, 39)


def test_week_53_compares_with_week_52():
    s = pd.Series({"2025-W52": 5.0})
    aligned = analysis.align_season(s, 2025, analysis.season_weeks(2026))
    assert aligned[analysis.season_weeks(2026).index((2026, 53))] == 5.0


def test_historical_comparison():
    result = analysis.analyse(make_raw())
    s = result["summary"]["NO"]["flu_cases"]
    assert result["latest_week"] == "2025-W41"
    assert s["value"] == 50 and s["prev"] == 40 and s["wow_pct"] == 25.0
    assert s["hist_mean"] == 20 and s["hist_min"] == 10 and s["hist_max"] == 30
    assert s["index"] == 250 and s["position"] == "over"
    assert s["season_to_date"] == 90 and s["season_to_date_mean"] == 40 and s["season_to_date_index"] == 225
    assert s["intensity"] == "Very high"  # 50 is above the mean season peak (20)
    assert result["summary"]["NO"]["flu_positivity"]["value"] == 50.0
    assert result["subtypes"]["NO"]["season"]["A(H3N2)"] == 7


def test_covid_and_hospital_indicators():
    result = analysis.analyse(make_raw())
    cov = result["summary"]["NO"]["covid_cases"]
    assert cov["value"] == 100 and cov["hist_mean"] == 200 and cov["index"] == 50
    hosp = result["summary"]["SE"]["flu_hospital"]
    assert hosp["value"] == 15 and hosp["index"] == 75
    assert any("COVID-19" in f for f in result["findings"])
    assert any("hospital admissions with influenza" in f for f in result["findings"])


def test_excel_workbook(tmp_path):
    from openpyxl import load_workbook
    from flutracker import excel

    path = excel.build(analysis.analyse(make_raw()), tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == excel.SHEET_ORDER
    ws = wb["Overview"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("ROUND(G" in f for f in formulas)  # index is a live formula
    assert ws["B2"].font.name == "Georgia"


def test_email_is_ascii_and_outlook_safe():
    html = emailer.render(analysis.analyse(make_raw()))
    assert html.isascii()
    assert "Respiratory virus update" in html
    assert "<div" not in html and "<style" not in html
    assert "v:roundrect" in html


def test_sales_sheets_and_import(tmp_path):
    from openpyxl import load_workbook
    from flutracker import excel, sales

    csv = tmp_path / "my_sales.csv"
    csv.write_text("country,week,sales\nNorway,2025-W40,1200\nSE,2025-W41,800\nDanmark,2025-W41,\n")
    rows = sales.read_sales(csv)
    assert rows == {("week", "2025-W40", "NO"): 1200.0, ("week", "2025-W41", "SE"): 800.0}

    path = excel.build(analysis.analyse(make_raw()), tmp_path / "out.xlsx", rows)
    wb = load_workbook(path)
    inp = wb["Sales input (weekly)"]
    found = {inp.cell(r, 2).value: (inp.cell(r, 4).value, inp.cell(r, 5).value) for r in range(sales.HR + 1, inp.max_row + 1)}
    assert found["2025-W40"] == (1200.0, None) and found["2025-W41"] == (None, 800.0)
    fit = wb["Sales vs flu (weekly)"]
    assert fit["D12"].value.startswith('=IF(D11<3,"",CORREL(')
    assert fit["D18"].value.startswith("=IF(COUNT(")  # forecast sum
    corr = wb["Correlation (monthly)"]
    assert any(isinstance(c.value, str) and "CORREL(" in c.value for row in corr.iter_rows() for c in row)


def test_month_assignment_uses_thursday():
    from flutracker import sales

    assert sales.month_of_week(2026, 1) == (2026, 1)   # Thu 1 Jan 2026
    assert sales.month_of_week(2025, 40) == (2025, 10)  # Thu 2 Oct 2025
    assert (2026, 1) in sales._weeks_in_month(2026, 1) and (2025, 53) not in sales._weeks_in_month(2026, 1)


def test_rsv_and_influenza_types():
    raw = make_raw()
    rsv = dict(pathogen="RSV", pathogentype="RSV", pathogensubtype="RSV", age="total", survtype="non-sentinel")
    typed = dict(pathogen="Influenza", age="total", survtype="non-sentinel", indicator="detections")
    data = []
    for start, base in [(2022, 10), (2023, 20), (2024, 30)]:
        data += rows("Norway", "nonSentinelTestsDetections", season_values(start, base), indicator="detections", **rsv)
        data += rows("Norway", "nonSentinelTestsDetections", season_values(start, 100), indicator="tests", **rsv)
    current = {"2025-W40": 10, "2025-W41": 40}
    data += rows("Norway", "nonSentinelTestsDetections", current, indicator="detections", **rsv)
    data += rows("Norway", "nonSentinelTestsDetections", {k: 200 for k in current}, indicator="tests", **rsv)
    # Influenza A is the sum of its subtypes; positivity is against all influenza tests (100 a week).
    for sub, v in [("A(H1)pdm09", 20), ("A(H3)", 10), ("A (unknown)", 5)]:
        data += rows("Norway", "nonSentinelTestsDetections", {"2025-W41": v}, pathogentype="Influenza A", pathogensubtype=sub, **typed)
    data += rows("Norway", "nonSentinelTestsDetections", {"2025-W41": 15}, pathogentype="Influenza B", pathogensubtype="B/Vic", **typed)
    result = analysis.analyse(pd.concat([raw, pd.DataFrame(data)], ignore_index=True))
    no = result["summary"]["NO"]
    assert no["rsv_cases"]["value"] == 40 and no["rsv_cases"]["hist_mean"] == 20 and no["rsv_cases"]["index"] == 200
    assert no["rsv_positivity"]["value"] == 20.0
    assert no["flu_a_cases"]["value"] == 35 and no["flu_a_positivity"]["value"] == 35.0
    assert no["flu_b_cases"]["value"] == 15 and no["flu_b_positivity"]["value"] == 15.0
    assert any(f.startswith("RSV, lab-confirmed") for f in result["findings"])
    assert any(f.startswith("Influenza A and B") and "A 35, B 15" in f for f in result["findings"])
    assert "RSV" in emailer.render(result)


def test_intensity_explained_with_example():
    result = analysis.analyse(make_raw())
    s = result["summary"]["NO"]["flu_cases"]
    # Peaks 10, 20 and 30 average 20; 50 / 20 = 250% of the average peak.
    assert s["peak_mean"] == 20 and s["intensity"] == "Very high"
    example = emailer.intensity_example(result)
    assert "peaks of 10, 20 and 30 average 20" in example and "250%" in example and "very high" in example


def test_recent_weeks_per_country_and_nordic():
    raw = make_raw()
    extra = rows("Sweden", "nonSentinelTestsDetections", {"2025-W41": 5}, indicator="detections", pathogen="SARS-CoV-2",
                 pathogentype="SARS-CoV-2", pathogensubtype="SARS-CoV-2", age="total", survtype="non-sentinel")
    result = analysis.analyse(pd.concat([raw, pd.DataFrame(extra)], ignore_index=True))
    recent = result["recent"]
    assert len(recent["weeks"]) == 52 and recent["weeks"][-1] == "2025-W41" and recent["weeks"][0] == "2024-W42"
    assert recent["countries"]["NO"]["covid_cases"][-1] == 100
    # Sweden reports only W41, so the Nordic sum exists for W41 and is left empty the week before.
    assert recent["nordic"]["covid_cases"][-1] == 105 and recent["nordic"]["covid_cases"][-2] is None


def test_influenza_type_missing_week_counts_as_zero():
    raw = make_raw()
    extra = rows("Norway", "nonSentinelTestsDetections", {"2025-W40": 3}, pathogen="Influenza", pathogentype="Influenza B",
                 pathogensubtype="B/Vic", indicator="detections", age="total", survtype="non-sentinel")
    result = analysis.analyse(pd.concat([raw, pd.DataFrame(extra)], ignore_index=True))
    b = result["summary"]["NO"]["flu_b_cases"]
    # Influenza was reported in W41 but no B row exists: that is 0 cases, not a missing week.
    assert b["value"] == 0 and b["prev"] == 3 and b["week"] == "2025-W41"
