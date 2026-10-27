"""Karo-branded Excel workbook with this week's figures and the full history.

Comparisons (change, index, level, historical average/min/max) are Excel
formulas over the weekly values, so the workbook recalculates if a value is
edited. Weekly values themselves come from ECDC ERVISS (see Sources sheet).
"""

import datetime as dt

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import sales
from .analysis import parse_yw, week_monday
from .config import COUNTRIES, INDICATORS, INTENSITY_LEVELS, RECENT_KEYS, SOURCE

NAVY = "110B36"
PEACH = "FAD7C2"
TEXT = "1A1A2E"
MUTED = "6B7280"
FONT = "Georgia"

F_BASE = Font(name=FONT, size=11, color=TEXT)
F_BOLD = Font(name=FONT, size=11, color=TEXT, bold=True)
F_HEAD = Font(name=FONT, size=11, color="FFFFFF", bold=True)
F_TITLE = Font(name=FONT, size=22, color="FFFFFF")
F_SUB = Font(name=FONT, size=11, color=PEACH)
F_LABEL = Font(name=FONT, size=9, color=MUTED, bold=True)
F_MUTED = Font(name=FONT, size=10, color=MUTED, italic=True)
F_LINK = Font(name=FONT, size=11, color=NAVY, underline="single")
FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_PEACH = PatternFill("solid", fgColor=PEACH)
FILL_PEACH_LIGHT = PatternFill("solid", fgColor="FDEFE6")
LINE = Border(bottom=Side(style="thin", color="E8E8E8"))
WRAP = Alignment(wrap_text=True, vertical="top")


def _banner(ws, title, subtitle, width):
    for r in (1, 2, 3):
        for c in range(1, width + 1):
            ws.cell(r, c).fill = FILL_NAVY
    ws.row_dimensions[1].height = 12
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[3].height = 24
    ws.cell(2, 2, title).font = F_TITLE
    ws.cell(3, 2, subtitle).font = F_SUB
    ws.cell(3, 2).alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 3
    ws.sheet_view.showGridLines = False


def _label(ws, row, col, text):
    c = ws.cell(row, col, text.upper())
    c.font = F_LABEL


def _header(ws, row, col, headers):
    for i, h in enumerate(headers):
        c = ws.cell(row, col + i, h)
        c.font, c.fill = F_HEAD, FILL_NAVY
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _number_format(key):
    meta = INDICATORS[key]
    return "#,##0" if meta["decimals"] == 0 else "#,##0.0"


def _week_label(result):
    start = dt.date.fromisoformat(result["latest_week_start"])
    end = start + dt.timedelta(days=6)
    return f"Week {parse_yw(result['latest_week'])[1]} · {start:%d %b} – {end:%d %b %Y} · Season {result['season']}"


def _intensity_formula(value, peak):
    """Nested IF that mirrors config.INTENSITY_LEVELS."""
    levels = INTENSITY_LEVELS
    expr = f'"{levels[-1][1]}"'
    for limit, name in reversed(levels[:-1]):
        expr = f'IF({value}/{peak}<{limit},"{name}",{expr})'
    return f'=IF(OR({value}="",{peak}="",{peak}=0),"",{expr})'


# ----------------------------------------------------------- overview ----

def _overview(wb, result):
    ws = wb.active
    ws.title = "Overview"
    headers = ["Group", "Country", "Indicator", "Unit", "Week", "This week", "Previous week", "Change",
               "Hist. average", "Hist. min", "Hist. max", "Index", "Level vs history",
               "Avg. season peak", "Intensity", "Season to date", "Hist. season to date", "Season-to-date index"]
    width = len(headers) + 1
    _banner(ws, "Flu, COVID-19 & RSV tracker – Norway, Sweden, Denmark", _week_label(result), width)

    r = 5
    _label(ws, r, 2, "Key findings this week")
    for f in result["findings"]:
        r += 1
        ws.cell(r, 2, "•").font = Font(name=FONT, size=11, color=NAVY, bold=True)
        ws.cell(r, 2).alignment = Alignment(horizontal="right", vertical="top")
        ws.cell(r, 3, f).font = F_BASE
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=width)
        ws.cell(r, 3).alignment = WRAP
        ws.row_dimensions[r].height = 16 * (len(f) // 160 + 1) + 2

    r += 2
    _label(ws, r, 2, "This week against the same week in previous seasons")
    r += 1
    top = r
    _header(ws, r, 2, headers)
    n = 0
    for group in result["groups"]:
        for code, c in COUNTRIES.items():
            for key, meta in INDICATORS.items():
                s = result["summary"].get(code, {}).get(key)
                if meta["group"] != group or not s:
                    continue
                r += 1
                n += 1
                fmt = _number_format(key)
                cells = [group, c["name"], meta["label"], meta["unit"], parse_yw(s["week"])[1],
                         s["value"], s["prev"],
                         f'=IF(OR(G{r}="",H{r}="",H{r}=0),"",(G{r}-H{r})/H{r})',
                         s["hist_mean"], s["hist_min"], s["hist_max"],
                         f'=IF(OR(J{r}="",J{r}=0),"",ROUND(G{r}/J{r}*100,0))',
                         f'=IF(J{r}="","",IF(G{r}>L{r},"Above hist. max",IF(G{r}<K{r},"Below hist. min","Within range")))',
                         s["peak_mean"], _intensity_formula(f"G{r}", f"O{r}"),
                         s.get("season_to_date"), s.get("season_to_date_mean"),
                         f'=IF(OR(Q{r}="",R{r}="",R{r}=0),"",ROUND(Q{r}/R{r}*100,0))']
                for i, v in enumerate(cells):
                    cell = ws.cell(r, 2 + i, v)
                    cell.font = F_BASE
                    cell.border = LINE
                    if n % 2 == 0:
                        cell.fill = FILL_PEACH_LIGHT
                for col in "GHJKLOQR":
                    ws[f"{col}{r}"].number_format = fmt
                ws[f"I{r}"].number_format = '+0%;-0%;0%'
                ws[f"M{r}"].font = F_BOLD
                ws[f"S{r}"].font = F_BOLD
                if s.get("low_tests"):
                    ws[f"D{r}"].value = meta["label"] + " †"
    ws.auto_filter.ref = f"B{top}:{get_column_letter(width)}{r}"
    ws.freeze_panes = ws.cell(top + 1, 5)
    r += 2
    for note in [
        "Index = this week / historical average for the same ISO week × 100 (100 = normal for the time of year).",
        "Intensity = this week / average of previous season peaks: <10% baseline, 10–40% low, 40–70% medium, 70–100% high, "
        ">100% very high. Column O holds the average peak; see 'Sources & method' for the full explanation.",
        "† fewer than 20 sentinel samples this week – interpret with care. Seasons run from week 40 to week 39.",
        f"Source: {SOURCE['name']}. {SOURCE['note']}",
    ]:
        ws.cell(r, 2, note).font = F_MUTED
        r += 1
    widths = [3, 20, 11, 44, 12, 7, 11, 11, 9, 11, 10, 10, 8, 17, 12, 11, 12, 13, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ------------------------------------------------------------- curves ----

def _curves(wb, result, key, title):
    ws = wb.create_sheet(title)
    meta = INDICATORS[key]
    weeks = result["season_weeks"]
    keys = result["season_week_keys"]
    codes = [c for c in COUNTRIES if key in result["series"].get(c, {})]
    if not codes:
        return
    # Countries can have different histories (e.g. RSV), so take every season any of them has.
    seasons = sorted({s for c in codes for s in result["series"][c][key]["seasons"]})
    prior, current = seasons[:-1], seasons[-1]
    block = len(seasons) + 4  # seasons + avg + min + max + index
    width = 2 + len(codes) * (block + 1)
    _banner(ws, f"{meta['label']} – season curves", _week_label(result) + f" · {meta['unit']}", width)

    hr = 6  # header row; country names in row 5
    ws.cell(hr, 2, "Week").font = F_HEAD
    ws.cell(hr, 2).fill = FILL_NAVY
    ws.cell(hr, 3, "Week starting").font = F_HEAD
    ws.cell(hr, 3).fill = FILL_NAVY
    for i, wk in enumerate(weeks):
        ws.cell(hr + 1 + i, 2, wk).font = F_BASE
        d = ws.cell(hr + 1 + i, 3, week_monday(*parse_yw(keys[i])))
        d.number_format = "dd.mm.yyyy"
        d.font = F_BASE
    last = hr + len(weeks)

    col = 4
    fmt = _number_format(key)
    for code in codes:
        data = result["series"][code][key]["seasons"]
        ws.cell(5, col, COUNTRIES[code]["name"]).font = Font(name=FONT, size=13, color=NAVY, bold=True)
        _header(ws, hr, col, seasons + ["Hist. average", "Hist. min", "Hist. max", "Index"])
        first, last_prior = get_column_letter(col), get_column_letter(col + len(prior) - 1)
        cur_col = get_column_letter(col + len(prior))
        avg_col = get_column_letter(col + len(seasons))
        for i in range(len(weeks)):
            r = hr + 1 + i
            for j, season in enumerate(seasons):
                v = data[season][i] if season in data else None
                cell = ws.cell(r, col + j, v)
                cell.number_format = fmt
                cell.font = F_BOLD if season == current else F_BASE
                if season == current:
                    cell.fill = FILL_PEACH
            rng = f"{first}{r}:{last_prior}{r}"
            ws.cell(r, col + len(seasons), f'=IF(COUNT({rng})<2,"",AVERAGE({rng}))')
            ws.cell(r, col + len(seasons) + 1, f'=IF(COUNT({rng})<2,"",MIN({rng}))')
            ws.cell(r, col + len(seasons) + 2, f'=IF(COUNT({rng})<2,"",MAX({rng}))')
            ws.cell(r, col + len(seasons) + 3,
                    f'=IF(OR({cur_col}{r}="",{avg_col}{r}="",{avg_col}{r}=0),"",ROUND({cur_col}{r}/{avg_col}{r}*100,0))')
            for k in range(4):
                c = ws.cell(r, col + len(seasons) + k)
                c.font = F_BASE
                c.number_format = "0" if k == 3 else fmt
            for k in range(block):
                ws.cell(r, col + k).border = LINE
        for k in range(block):
            ws.column_dimensions[get_column_letter(col + k)].width = 11

        chart = LineChart()
        chart.title = f"{COUNTRIES[code]['name']} – {meta['short']}"
        chart.height, chart.width = 7.5, 16
        chart.y_axis.title = meta["unit"]
        chart.legend.position = "b"
        cats = Reference(ws, min_col=2, min_row=hr + 1, max_row=last)
        for c_idx, colour, w, dash in [
            (col + len(seasons) + 2, "E9B79A", 12700, None),          # hist max
            (col + len(seasons) + 1, "E9B79A", 12700, None),          # hist min
            (col + len(seasons), "6B7280", 19050, "dash"),            # hist average
            (col + len(prior), NAVY, 34925, None),                    # current season
        ]:
            ref = Reference(ws, min_col=c_idx, min_row=hr, max_row=last)
            chart.add_data(ref, titles_from_data=True)
            series = chart.series[-1]
            series.graphicalProperties.line.solidFill = colour
            series.graphicalProperties.line.width = w
            if dash:
                series.graphicalProperties.line.dashStyle = dash
            series.smooth = False
        chart.set_categories(cats)
        ws.add_chart(chart, f"{get_column_letter(col)}{last + 3}")
        col += block + 1

    ws.cell(4, 2, "Each block: weekly values per season (current season shaded), then historical average, min, max and index "
                  "calculated with formulas from the previous seasons.").font = F_MUTED
    ws.freeze_panes = ws.cell(hr + 1, 4)
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 13


# ------------------------------------------------------ last 12 months ----

RECENT_STYLE = {  # colour and dash per virus, as on the dashboard
    "flu_a_cases": ("Influenza A", NAVY, "solid"),
    "flu_b_cases": ("Influenza B", "9B3F96", "sysDot"),
    "covid_cases": ("COVID-19", "E07B39", "lgDash"),
    "rsv_cases": ("RSV", "008A7A", "dash"),
}


def _recent(wb, result):
    rec = result["recent"]
    areas = [("Nordics", rec["nordic"])] + [(COUNTRIES[c]["name"], v) for c, v in rec["countries"].items()]
    ws = wb.create_sheet("Last 12 months")
    width = 3 + len(areas) * len(RECENT_KEYS)
    _banner(ws, "Lab-confirmed cases by virus – last 12 months", _week_label(result), width)
    ws.cell(4, 2, "Weekly lab-confirmed cases per virus. Blank = not reported to ECDC that week. The Nordic total is "
                  "only given for weeks in which all three countries reported.").font = F_MUTED

    hr = 40  # charts above, table below
    _header(ws, hr, 2, ["Week", "Week starting"] + [f"{name} – {RECENT_STYLE[k][0]}" for name, _ in areas for k in RECENT_KEYS])
    for i, wk in enumerate(rec["weeks"]):
        r = hr + 1 + i
        ws.cell(r, 2, wk.replace("-W", "-")).font = F_BASE
        d = ws.cell(r, 3, week_monday(*parse_yw(wk)))
        d.number_format, d.font = "dd.mm.yyyy", F_BASE
        for a, (_, values) in enumerate(areas):
            for j, k in enumerate(RECENT_KEYS):
                v = values.get(k, [None] * len(rec["weeks"]))[i]
                cell = ws.cell(r, 4 + a * len(RECENT_KEYS) + j, v)
                cell.number_format, cell.font = "#,##0", F_BASE
        for c in range(2, width + 1):
            ws.cell(r, c).border = LINE
    last = hr + len(rec["weeks"])

    cats = Reference(ws, min_col=2, min_row=hr + 1, max_row=last)
    for a, (name, values) in enumerate(areas):
        chart = LineChart()
        chart.title = f"{name} – lab-confirmed cases"
        chart.height, chart.width = 8, 16
        chart.y_axis.title = "cases"
        chart.legend.position = "b"
        chart.display_blanks = "gap"
        for j, k in enumerate(RECENT_KEYS):
            if k not in values:
                continue
            col = 4 + a * len(RECENT_KEYS) + j
            chart.add_data(Reference(ws, min_col=col, min_row=hr, max_row=last), titles_from_data=True)
            series = chart.series[-1]
            series.tx = SeriesLabel(v=RECENT_STYLE[k][0])
            series.graphicalProperties.line.solidFill = RECENT_STYLE[k][1]
            series.graphicalProperties.line.width = 28575
            series.graphicalProperties.line.dashStyle = RECENT_STYLE[k][2]
            series.smooth = False
        chart.set_categories(cats)
        ws.add_chart(chart, f"{'B' if a % 2 == 0 else 'K'}{6 + (a // 2) * 17}")
    ws.freeze_panes = ws.cell(hr + 1, 4)
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 13
    for c in range(4, width + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12


# --------------------------------------------------------- weekly data ----

def _weekly_data(wb, result):
    ws = wb.create_sheet("Weekly data")
    headers = ["Country", "Group", "Indicator", "Unit", "Season", "ISO week", "Week starting", "Value"]
    _banner(ws, "Weekly data – all indicators and seasons", _week_label(result), len(headers) + 1)
    hr = 5
    for i, h in enumerate(headers):
        ws.cell(hr, 2 + i, h)
    r = hr
    for code, by_key in result["series"].items():
        for key, s in by_key.items():
            meta = INDICATORS[key]
            for season, values in s["seasons"].items():
                start_year = int(season[:4])
                for i, v in enumerate(values):
                    if v is None:
                        continue
                    wk = parse_yw(result["season_week_keys"][i])[1]
                    year = start_year if wk >= 40 else start_year + 1
                    if wk == 53 and dt.date(year, 12, 28).isocalendar()[1] != 53:
                        continue  # comparison slot only, not a real week in that season
                    r += 1
                    row = [COUNTRIES[code]["name"], meta["group"], meta["label"], meta["unit"], season,
                           f"{year}-W{wk:02d}", week_monday(year, wk), v]
                    for j, val in enumerate(row):
                        ws.cell(r, 2 + j, val).font = F_BASE
                    ws.cell(r, 8).number_format = "dd.mm.yyyy"
                    ws.cell(r, 9).number_format = _number_format(key)
    table = Table(displayName="WeeklyData", ref=f"B{hr}:{get_column_letter(1 + len(headers))}{r}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(table)
    _header(ws, hr, 2, headers)
    ws.freeze_panes = ws.cell(hr + 1, 2)
    for col, w in zip("ABCDEFGHI", [3, 12, 20, 46, 12, 10, 11, 14, 11]):
        ws.column_dimensions[col].width = w


# ------------------------------------------------------------ subtypes ----

def _subtypes(wb, result):
    ws = wb.create_sheet("Virus types")
    codes = [c for c in COUNTRIES if c in result["subtypes"]]
    width = 2 + 2 * len(codes) + 2
    _banner(ws, "Influenza virus types", _week_label(result), width)
    names = list(result["subtypes"][codes[0]]["season"])
    for b, (period, title) in enumerate([("week", "This week"), ("season", "Season to date")]):
        top = 6 + b * (len(names) + 6)
        _label(ws, top - 1, 2, f"{title} – lab-confirmed influenza by type")
        _header(ws, top, 2, ["Virus"] + [COUNTRIES[c]["name"] for c in codes] + ["All three"]
                + [f"Share {COUNTRIES[c]['name']}" for c in codes])
        for i, name in enumerate(names):
            r = top + 1 + i
            ws.cell(r, 2, name).font = F_BASE
            for j, c in enumerate(codes):
                ws.cell(r, 3 + j, result["subtypes"][c][period][name]).font = F_BASE
            first, lastc = get_column_letter(3), get_column_letter(2 + len(codes))
            ws.cell(r, 3 + len(codes), f"=SUM({first}{r}:{lastc}{r})").font = F_BOLD
            for j in range(len(codes)):
                colL = get_column_letter(3 + j)
                tot = f"{colL}{top + 1 + len(names)}"
                cell = ws.cell(r, 4 + len(codes) + j, f'=IF({tot}=0,"",{colL}{r}/{tot})')
                cell.number_format = "0%"
                cell.font = F_BASE
            if i % 2:
                for k in range(2, width):
                    ws.cell(r, k).fill = FILL_PEACH_LIGHT
        tr = top + 1 + len(names)
        ws.cell(tr, 2, "Total").font = F_BOLD
        for j in range(len(codes) + 1):
            colL = get_column_letter(3 + j)
            ws.cell(tr, 3 + j, f"=SUM({colL}{top + 1}:{colL}{tr - 1})").font = F_BOLD
    ws.column_dimensions["B"].width = 22
    for j in range(3, width + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14


# ------------------------------------------------------------- sources ----

def _sources(wb, result):
    ws = wb.create_sheet("Sources & method")
    _banner(ws, "Sources & method", _week_label(result), 6)
    r = 5
    _label(ws, r, 2, "Data source")
    rows = [
        ("Dataset", SOURCE["name"], SOURCE["url"]),
        ("Publisher", SOURCE["publisher"], None),
        ("Open data", "ECDC respiratory virus weekly data on GitHub", SOURCE["data"]),
        ("Reporting", SOURCE["note"], None),
    ]
    for code, c in COUNTRIES.items():
        rows.append((f"{c['name']} – reports", c["authority"], c["url"]))
        rows.append((f"{c['name']} – statistics", c["authority"] + " statistics", c["data_url"]))
    for label, text, url in rows:
        r += 1
        ws.cell(r, 2, label).font = F_BOLD
        cell = ws.cell(r, 3, text)
        cell.font, cell.alignment = F_BASE, WRAP
        if url:
            cell.hyperlink = url
            cell.font = F_LINK
    r += 2
    _label(ws, r, 2, "Method")
    for line in [
        "Seasons run from ISO week 40 to week 39 the following year.",
        "Each week is compared with the same ISO week in every previous season with at least 20 weeks of data "
        "(week 53 is compared with week 52).",
        "Index = this week / historical average × 100. 100 = normal for the time of year.",
        "Level: above the highest, below the lowest, or within the range of previous seasons for that week.",
        "Intensity – how big this week is compared with the top of a typical season. Calculated per country and indicator:",
        "   1. For each previous season (week 40–39) with at least 20 weeks of data, take the highest weekly value: that season's peak.",
        "   2. Average those peaks ('Avg. season peak' on the Overview sheet).",
        "   3. Divide this week's value by the average peak.",
        "   4. Under 10% = baseline, 10% to under 40% = low, 40% to under 70% = medium, 70% to under 100% = high, "
        "100% or more = very high.",
        "   Example: 171 cases this week and season peaks of 1,800, 2,400 and 1,950 give an average peak of 2,050; "
        "171 / 2,050 = 8% → baseline.",
        "   Intensity differs from the index: the index compares with the same week in earlier seasons (normal for the time "
        "of year?), intensity with the season peaks (how close to a full epidemic?). Fewer previous seasons make the "
        "average peak less reliable (RSV in Norway rests on two). ECDC and the national institutes use their own methods, "
        "so their intensity labels can differ.",
        "Season to date = sum since week 40, compared with previous seasons summed to the same week.",
        "Indicators a country has not reported for more than 8 weeks are left out. Hospital admissions are only "
        "included where the national institute reports them to ECDC.",
        f"Generated {result['generated_at'][:10]} from data up to {result['latest_week']}.",
    ]:
        r += 1
        ws.cell(r, 2, "•").font = Font(name=FONT, color=NAVY, bold=True)
        cell = ws.cell(r, 3, line)
        cell.font, cell.alignment = F_BASE, WRAP
        ws.row_dimensions[r].height = 30 if len(line) > 100 else 16
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 110


SHEET_ORDER = [
    "Overview", "Last 12 months", "Influenza curves", "Influenza A curves", "Influenza B curves", "COVID-19 curves", "RSV curves",
    "Sales input (weekly)", "Sales vs flu (weekly)", "Correlation (weekly)",
    "Sales input (monthly)", "Sales vs flu (monthly)", "Correlation (monthly)",
    "Weekly data", "Virus types", "Flu by week", "Flu by month", "Sources & method",
]


def build(result, path, sales_rows=None):
    wb = Workbook()
    _overview(wb, result)
    _recent(wb, result)
    _curves(wb, result, "flu_cases", "Influenza curves")
    _curves(wb, result, "flu_a_cases", "Influenza A curves")
    _curves(wb, result, "flu_b_cases", "Influenza B curves")
    _curves(wb, result, "covid_cases", "COVID-19 curves")
    _curves(wb, result, "rsv_cases", "RSV curves")
    _weekly_data(wb, result)
    _subtypes(wb, result)
    _sources(wb, result)
    for freq in ("week", "month"):
        sales.build(wb, result, freq)
    if sales_rows:
        sales.prefill(wb, sales_rows)
    wb._sheets.sort(key=lambda ws: SHEET_ORDER.index(ws.title))
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title == "Overview" else PEACH
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
