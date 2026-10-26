"""Sales sheets in the Excel workbook: input, correlation with flu, and forecast.

For each frequency (weekly and monthly) the workbook gets four sheets whose
data rows line up row for row (row r is the same period on every sheet):

  Sales input (…)    yellow cells where the user pastes sales per country
  Flu by week/month  the flu/COVID drivers per country; future periods hold the
                     historical average for that week × a scenario percentage
  Sales vs flu (…)   chosen driver + lag per country → correlation, R², a linear
                     fit (sales = intercept + slope × driver) and a forecast
  Correlation (…)    correlation of sales with every driver at every lag

Everything after the input is an Excel formula, so the workbook recalculates
as soon as sales are pasted in. Sales data never leaves the user's computer.
"""

import datetime as dt

from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .analysis import parse_yw, week_monday, weeks_in_year, yw_str
from .config import COUNTRIES, INDICATORS

DRIVERS = ["flu_cases", "flu_positivity", "flu_sentinel", "ili_rate", "covid_cases", "covid_positivity"]
DEFAULT_DRIVER = "flu_cases"
HR = 23  # header row of the row-aligned tables on every sales sheet
FILL_INPUT = PatternFill("solid", fgColor="FFF4CC")

FREQ = {
    "week": {"name": "weekly", "unit": "week", "units": "weeks", "history_start": (2022, 30),
             "ahead": 12, "max_lag": 8, "default_lag": 1, "sums": (4, 12)},
    "month": {"name": "monthly", "unit": "month", "units": "months", "history_start": (2022, 7),
              "ahead": 3, "max_lag": 3, "default_lag": 0, "sums": (1, 3)},
}


# --------------------------------------------------------------- periods ----

def weekly_values(result, code, key):
    """{'YYYY-Www': value} for every reported week of one country's indicator."""
    s = result["series"].get(code, {}).get(key)
    if not s:
        return {}
    out = {}
    for season, values in s["seasons"].items():
        start = int(season[:4])
        for i, v in enumerate(values):
            if v is None:
                continue
            wk = parse_yw(result["season_week_keys"][i])[1]
            year = start if wk >= 40 else start + 1
            if wk == 53 and weeks_in_year(year) != 53:
                continue  # comparison slot only, not a real week of that season
            out[yw_str(year, wk)] = v
    return out


def expected_weekly(result, code, key, yw):
    """Historical average for this ISO week (the forecast baseline), or None."""
    s = result["series"].get(code, {}).get(key)
    if not s or yw not in result["season_week_keys"]:
        return None
    return s["band"]["mean"][result["season_week_keys"].index(yw)]


def _next_week(y, w):
    return (y, w + 1) if w < weeks_in_year(y) else (y + 1, 1)


def week_periods(result):
    latest = parse_yw(result["latest_week"])
    cur, out = FREQ["week"]["history_start"], []
    while cur <= latest:
        out.append(cur)
        cur = _next_week(*cur)
    future = []
    for _ in range(FREQ["week"]["ahead"]):
        future.append(cur)
        cur = _next_week(*cur)
    return out, future


def month_of_week(y, w):
    """ISO convention: a week belongs to the month of its Thursday."""
    thu = week_monday(y, w) + dt.timedelta(days=3)
    return thu.year, thu.month


def month_periods(result):
    latest = parse_yw(result["latest_week"])
    last_month = month_of_week(*latest)
    y, m = FREQ["month"]["history_start"]
    out = []
    while (y, m) <= last_month:
        out.append((y, m))
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)
    future = []
    for _ in range(FREQ["month"]["ahead"]):
        future.append((y, m))
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)
    return out, future


def _weeks_in_month(y, m):
    first = dt.date(y, m, 1)
    d = first - dt.timedelta(days=first.weekday())  # Monday on/before the 1st
    weeks = []
    while d < first + dt.timedelta(days=40):
        iy, iw, _ = d.isocalendar()
        if month_of_week(iy, iw) == (y, m):
            weeks.append((iy, iw))
        d += dt.timedelta(days=7)
    return weeks


def monthly_value(result, code, key, ym, weekly):
    """(value, complete) for one month: sum of weeks for counts, mean for rates.

    A month with weeks after the latest data week is filled with the historical
    average for those weeks and reported as not complete."""
    latest = result["latest_week"]
    vals, complete = [], True
    for y, w in _weeks_in_month(*ym):
        k = yw_str(y, w)
        if k <= latest:
            v = weekly.get(k)
        else:
            complete = False
            v = expected_weekly(result, code, key, k)
        if v is not None:
            vals.append(v)
    if not vals:
        return None, complete
    total = sum(vals) if INDICATORS[key]["kind"] == "count" else sum(vals) / len(vals)
    return round(total, INDICATORS[key]["decimals"]), complete


# ---------------------------------------------------------------- sheets ----

def _sheet_names(freq):
    n = FREQ[freq]["name"]
    return {
        "input": f"Sales input ({n})",
        "fit": f"Sales vs flu ({n})",
        "corr": f"Correlation ({n})",
        "flu": "Flu by week" if freq == "week" else "Flu by month",
    }


def _flu_header(code, key):
    return f"{COUNTRIES[code]['name']} - {INDICATORS[key]['label']}"


def _period_cells(freq, p):
    if freq == "week":
        return yw_str(*p), week_monday(*p)
    return f"{p[0]}-{p[1]:02d}", dt.date(p[0], p[1], 1)


def build(wb, result, freq):
    """Add the four sales sheets for `freq` ('week' or 'month')."""
    from . import excel as x  # shared Karo styles; imported here to avoid a circular import

    cfg = FREQ[freq]
    names = _sheet_names(freq)
    hist, future = week_periods(result) if freq == "week" else month_periods(result)
    periods = hist + future
    codes = list(COUNTRIES)
    first, last = HR + 1, HR + len(periods)
    last_hist = HR + len(hist)
    # Statistics use complete periods only (the current month is still running).
    partial = freq == "month" and any(yw_str(*w) > result["latest_week"] for w in _weeks_in_month(*hist[-1]))
    stat_end = last_hist - 1 if partial else last_hist
    sub = x._week_label(result)

    # ---- flu drivers -------------------------------------------------------
    flu = wb.create_sheet(names["flu"])
    cols = [(c, k) for c in codes for k in DRIVERS]
    flu_width = 4 + len(cols)
    x._banner(flu, f"Flu & COVID-19 drivers by {cfg['unit']}", sub, flu_width)
    flu.cell(5, 2, f"Used by the sales sheets. Reported {cfg['units']} come from ECDC ERVISS. Forecast {cfg['units']} hold "
                   f"the historical average for that {cfg['unit']} × the flu scenario set on '{names['fit']}'.").font = x.F_MUTED
    if freq == "month":
        flu.cell(6, 2, "Monthly values: weeks are assigned to the month of their Thursday; counts are summed, rates averaged. "
                       "The current month is completed with historical averages.").font = x.F_MUTED
    x._header(flu, HR, 2, [cfg["unit"].title(), "Starts", "Status"] + [_flu_header(c, k) for c, k in cols])
    weekly = {(c, k): weekly_values(result, c, k) for c, k in cols}
    fit_scen = {c: f"'{names['fit']}'!${get_column_letter(4 + i * 5)}$10" for i, c in enumerate(codes)}
    for i, p in enumerate(periods):
        r = first + i
        key, start = _period_cells(freq, p)
        flu.cell(r, 2, key).font = x.F_BASE
        flu.cell(r, 3, start).number_format = "dd.mm.yyyy"
        for j, (c, k) in enumerate(cols):
            col = 5 + j
            fmt = x._number_format(k)
            if freq == "week":
                if p in future:
                    exp = expected_weekly(result, c, k, key)
                    v = None if exp is None else f"=ROUND({exp}*{fit_scen[c]},{INDICATORS[k]['decimals']})"
                    status = "Forecast"
                else:
                    v, status = weekly[(c, k)].get(key), "Reported"
            else:
                v, complete = monthly_value(result, c, k, p, weekly[(c, k)])
                if p in future:
                    v = None if v is None else f"=ROUND({v}*{fit_scen[c]},{INDICATORS[k]['decimals']})"
                    status = "Forecast"
                else:
                    status = "Reported" if complete else "Partial"
            cell = flu.cell(r, col, v)
            cell.number_format, cell.font = fmt, x.F_BASE
        flu.cell(r, 4, status).font = x.F_BOLD if status != "Reported" else x.F_BASE
        if status != "Reported":
            for col in range(2, flu_width + 1):
                flu.cell(r, col).fill = x.FILL_PEACH_LIGHT
    for col in range(5, flu_width + 1):
        flu.column_dimensions[get_column_letter(col)].width = 16
    flu.column_dimensions["B"].width, flu.column_dimensions["C"].width = 11, 12
    flu.freeze_panes = flu.cell(first, 5)
    # Driver names for the drop-downs.
    flu.cell(8, 2, "DRIVER LIST").font = x.F_LABEL
    for i, k in enumerate(DRIVERS):
        flu.cell(9 + i, 2, INDICATORS[k]["label"]).font = x.F_BASE
    driver_list = f"'{names['flu']}'!$B$9:$B${8 + len(DRIVERS)}"
    flu_matrix = f"'{names['flu']}'!$E${first}:${get_column_letter(flu_width)}${last}"
    flu_heads = f"'{names['flu']}'!$E${HR}:${get_column_letter(flu_width)}${HR}"

    # ---- sales input -------------------------------------------------------
    inp = wb.create_sheet(names["input"])
    x._banner(inp, f"Sales input – {cfg['name']}", sub, 6)
    x._label(inp, 5, 2, "How to use")
    steps = [
        f"1. Paste {cfg['name']} sales per country into the yellow cells (units or NOK – use the same measure throughout).",
        f"2. Leave a cell empty when there is no figure. At least 12 {cfg['units']} of sales are needed for a meaningful correlation.",
        f"3. Open '{names['corr']}' to see which flu/COVID indicator and lag fits your sales best.",
        f"4. On '{names['fit']}', pick the driver and lag (yellow cells) and read the correlation and the forecast "
        f"for the next {cfg['ahead']} {cfg['units']}.",
        "Your sales stay in this file only – nothing is uploaded. Correlation shows co-movement, not cause and effect.",
    ]
    for i, s in enumerate(steps):
        inp.cell(6 + i, 2, s).font = x.F_BASE if i < 4 else x.F_MUTED
    x._label(inp, 13, 2, "Example of the expected format")
    ex_key = "2025-W45" if freq == "week" else "2025-11"
    for j, v in enumerate([ex_key, "", 1250, 830, 640]):
        c = inp.cell(14, 2 + j, v)
        c.font = x.F_MUTED
    inp.cell(15, 2, "(the example row above is not used in any calculation)").font = x.F_MUTED
    x._header(inp, HR, 2, [cfg["unit"].title(), "Starts"] + [f"{COUNTRIES[c]['name']} sales" for c in codes])
    for i, p in enumerate(periods):
        r = first + i
        key, start = _period_cells(freq, p)
        inp.cell(r, 2, key).font = x.F_BASE
        inp.cell(r, 3, start).number_format = "dd.mm.yyyy"
        for j in range(len(codes)):
            c = inp.cell(r, 4 + j)
            c.fill, c.number_format, c.font = FILL_INPUT, "#,##0", x.F_BASE
    for col, w in zip("ABCDEF", [3, 12, 12, 16, 16, 16]):
        inp.column_dimensions[col].width = w
    inp.freeze_panes = inp.cell(first, 4)

    # ---- correlation grid --------------------------------------------------
    corr = wb.create_sheet(names["corr"])
    lags = list(range(cfg["max_lag"] + 1))
    width = 2 + 1 + len(lags) + 2
    x._banner(corr, f"Correlation of sales with flu – {cfg['name']}", sub, width)
    corr.cell(5, 2, f"Pearson correlation (r) between sales in a {cfg['unit']} and each indicator L {cfg['units']} earlier. "
                    "+1 = move together, 0 = no relation, −1 = opposite. Uses reported periods with sales only.").font = x.F_MUTED
    start_row = first + cfg["max_lag"]  # same period range for every lag
    best = {}
    r = 7
    for ci, c in enumerate(codes):
        sales_col = get_column_letter(4 + ci)
        corr.cell(r, 2, COUNTRIES[c]["name"]).font = Font(name=x.FONT, size=13, color=x.NAVY, bold=True)
        r += 1
        x._header(corr, r, 2, ["Indicator"] + [f"Lag {L}" for L in lags] + ["Best r", "Best lag"])
        top = r + 1
        for di, k in enumerate(DRIVERS):
            r += 1
            corr.cell(r, 2, INDICATORS[k]["label"]).font = x.F_BASE
            fcol = get_column_letter(5 + codes.index(c) * len(DRIVERS) + di)
            for li, L in enumerate(lags):
                sales = f"'{names['input']}'!{sales_col}{start_row}:{sales_col}{stat_end}"
                drv = f"'{names['flu']}'!{fcol}{start_row - L}:{fcol}{stat_end - L}"
                cell = corr.cell(r, 3 + li, f"=IFERROR(CORREL({sales},{drv}),\"\")")
                cell.number_format, cell.font = "0.00", x.F_BASE
            row_rng = f"C{r}:{get_column_letter(2 + len(lags))}{r}"
            bc, bl = get_column_letter(3 + len(lags)), get_column_letter(4 + len(lags))
            corr[f"{bc}{r}"] = f'=IF(COUNT({row_rng})=0,"",MAX({row_rng}))'
            corr[f"{bl}{r}"] = (f'=IF({bc}{r}="","",INDEX($C${top - 1}:${get_column_letter(2 + len(lags))}${top - 1},'
                                f'MATCH({bc}{r},{row_rng},0)))')
            corr[f"{bc}{r}"].number_format = "0.00"
            corr[f"{bc}{r}"].font = corr[f"{bl}{r}"].font = x.F_BOLD
        bottom = r
        grid = f"C{top}:{get_column_letter(2 + len(lags))}{bottom}"
        corr.conditional_formatting.add(grid, ColorScaleRule(start_type="num", start_value=-1, start_color="FFFFFF",
                                                             mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                             end_type="num", end_value=1, end_color=x.PEACH))
        bc = get_column_letter(3 + len(lags))
        r += 1
        best_rng = f"${bc}${top}:${bc}${bottom}"
        corr.cell(r, 2, "Best combination").font = x.F_BOLD
        corr.cell(r, 3, f'=IF(COUNT({best_rng})=0,"Add sales to see the best match",'
                        f'INDEX($B${top}:$B${bottom},MATCH(MAX({best_rng}),{best_rng},0))&", "&'
                        f'INDEX(${get_column_letter(4 + len(lags))}${top}:${get_column_letter(4 + len(lags))}${bottom},'
                        f'MATCH(MAX({best_rng}),{best_rng},0))&" (r = "&TEXT(MAX({best_rng}),"0.00")&")")').font = x.F_BOLD
        best[c] = f"'{names['corr']}'!$C${r}"
        r += 2
    corr.column_dimensions["B"].width = 44
    for li in range(len(lags) + 2):
        corr.column_dimensions[get_column_letter(3 + li)].width = 9

    # ---- fit & forecast ----------------------------------------------------
    fit = wb.create_sheet(names["fit"])
    per = 5
    fit_width = 3 + per * len(codes)
    x._banner(fit, f"Sales vs flu – correlation & forecast ({cfg['name']})", sub, fit_width)
    fit.cell(5, 2, f"Linear model per country: sales = intercept + slope × driver, where the driver is the chosen indicator "
                   f"L {cfg['units']} earlier. Fitted on reported {cfg['units']} with sales; the forecast uses reported flu "
                   f"where the lag allows, otherwise the historical average × the flu scenario.").font = x.F_MUTED
    x._header(fit, 7, 2, ["Setting / result"])
    labels = ["Driver (choose)", f"Lag in {cfg['units']} (choose)", "Flu scenario for the forecast (% of a normal season)",
              f"{cfg['units'].title()} with sales and driver (n)", "Correlation (r)", "R² (share of variation explained)",
              "Slope (sales per unit of driver)", "Intercept", "Standard error of a forecast",
              "Strength of relationship", f"Forecast, next {cfg['sums'][0]} {cfg['unit'] if cfg['sums'][0] == 1 else cfg['units']}",
              f"Forecast, next {cfg['sums'][1]} {cfg['units']}", "Best match on the correlation sheet"]
    for i, lab in enumerate(labels):
        fit.cell(8 + i, 2, lab).font = x.F_BOLD if i < 3 else x.F_BASE
    dv_driver = DataValidation(type="list", formula1=f"={driver_list}", allow_blank=False)
    dv_lag = DataValidation(type="whole", operator="between", formula1="0", formula2=str(cfg["max_lag"]),
                            error=f"Choose a lag between 0 and {cfg['max_lag']}", showErrorMessage=True)
    fit.add_data_validation(dv_driver)
    fit.add_data_validation(dv_lag)

    n_fut = len(future)
    for ci, c in enumerate(codes):
        base = 4 + ci * per  # first table column for this country
        S, D, F, Lo, Hi = (get_column_letter(base + k) for k in range(per))
        sc = S  # settings and results sit above the country's own table block
        x._header(fit, 7, base, [COUNTRIES[c]["name"]])
        for rr in range(7, 21):
            fit.merge_cells(f"{S}{rr}:{Hi}{rr}")
        sales_h = f"${S}${first}:${S}${stat_end}"
        drv_h = f"${D}${first}:${D}${stat_end}"
        fit[f"{sc}8"] = INDICATORS[DEFAULT_DRIVER]["label"]
        fit[f"{sc}9"] = cfg["default_lag"]
        fit[f"{sc}10"] = 1
        for rr in (8, 9, 10):
            fit[f"{sc}{rr}"].fill = FILL_INPUT
        dv_driver.add(f"{sc}8")
        dv_lag.add(f"{sc}9")
        fit[f"{sc}10"].number_format = "0%"
        fit[f"{sc}11"] = f"=SUMPRODUCT(--ISNUMBER({sales_h}),--ISNUMBER({drv_h}))"
        fit[f"{sc}12"] = f'=IF({sc}11<3,"",CORREL({sales_h},{drv_h}))'
        fit[f"{sc}13"] = f'=IF({sc}12="","",RSQ({sales_h},{drv_h}))'
        fit[f"{sc}14"] = f'=IF({sc}12="","",SLOPE({sales_h},{drv_h}))'
        fit[f"{sc}15"] = f'=IF({sc}12="","",INTERCEPT({sales_h},{drv_h}))'
        fit[f"{sc}16"] = f'=IF({sc}11<4,"",STEYX({sales_h},{drv_h}))'
        fit[f"{sc}17"] = (f'=IF({sc}11<12,"Too little data (need 12+ {cfg["units"]})",IF(ABS({sc}12)>=0.7,"Strong",'
                          f'IF(ABS({sc}12)>=0.4,"Moderate",IF(ABS({sc}12)>=0.2,"Weak","No clear")))&'
                          f'IF(ABS({sc}12)>=0.2,IF({sc}12>0," positive"," negative"),""))')
        f1, f2 = last_hist + 1, last_hist + cfg["sums"][0]
        fit[f"{sc}18"] = f'=IF(COUNT(${F}${f1}:${F}${f2})=0,"",SUM(${F}${f1}:${F}${f2}))'
        f2 = last_hist + min(cfg["sums"][1], n_fut)
        fit[f"{sc}19"] = f'=IF(COUNT(${F}${f1}:${F}${f2})=0,"",SUM(${F}${f1}:${F}${f2}))'
        fit[f"{sc}20"] = f"={best[c]}"
        for rr, fmt in [(11, "0"), (12, "0.00"), (13, "0.00"), (14, "#,##0.00"), (15, "#,##0"), (16, "#,##0"),
                        (18, "#,##0"), (19, "#,##0")]:
            fit[f"{sc}{rr}"].number_format = fmt
        for rr in range(8, 21):
            fit[f"{sc}{rr}"].font = x.F_BOLD if rr in (12, 17, 18, 19) else x.F_BASE
            fit[f"{sc}{rr}"].alignment = Alignment(horizontal="right" if rr not in (8, 17, 20) else "left", wrap_text=rr == 20)
            fit[f"{sc}{rr}"].border = x.LINE

        x._header(fit, HR, base, [f"{COUNTRIES[c]['name']} sales", "Driver (lagged)", "Fitted / forecast",
                                  "Low (95%)", "High (95%)"])
        head = f'{sc}$8'
        match = f'MATCH("{COUNTRIES[c]["name"]} - "&${head},{flu_heads},0)'
        for i in range(len(periods)):
            r = first + i
            n = i + 1
            fit[f"{S}{r}"] = f"=IF('{names['input']}'!{get_column_letter(4 + ci)}{r}=\"\",\"\",'{names['input']}'!{get_column_letter(4 + ci)}{r})"
            pick = f"INDEX({flu_matrix},{n}-${sc}$9,{match})"
            fit[f"{D}{r}"] = f'=IFERROR(IF({n}-${sc}$9<1,"",IF({pick}="","",{pick})),"")'
            fit[f"{F}{r}"] = f'=IF(OR({D}{r}="",${sc}$14=""),"",MAX(0,${sc}$15+${sc}$14*{D}{r}))'
            if r > last_hist:
                fit[f"{Lo}{r}"] = f'=IF(OR({F}{r}="",${sc}$16=""),"",MAX(0,{F}{r}-1.96*${sc}$16))'
                fit[f"{Hi}{r}"] = f'=IF(OR({F}{r}="",${sc}$16=""),"",{F}{r}+1.96*${sc}$16)'
            for col in (S, F, Lo, Hi):
                fit[f"{col}{r}"].number_format = "#,##0"
            fit[f"{D}{r}"].number_format = "#,##0.0"
            for col in (S, D, F, Lo, Hi):
                fit[f"{col}{r}"].font = x.F_BOLD if (r > last_hist and col == F) else x.F_BASE
                fit[f"{col}{r}"].border = x.LINE
                if r > last_hist:
                    fit[f"{col}{r}"].fill = x.FILL_PEACH_LIGHT
        for k in range(per):
            fit.column_dimensions[get_column_letter(base + k)].width = 13

        chart = LineChart()
        chart.title = f"{COUNTRIES[c]['name']} – sales, fitted and forecast"
        chart.height, chart.width = 7.5, 18
        chart.legend.position = "b"
        chart.display_blanks = "gap"
        # Sales are charted from the input sheet, where missing weeks are truly empty and show as gaps
        # (formula cells returning "" would be drawn as zero). Charts start after the longest lag.
        c0 = first + cfg["max_lag"]
        for sheet, col, colour, w in [(inp, 4 + ci, x.NAVY, 28575), (fit, base + 2, "C8603A", 19050)]:
            chart.add_data(Reference(sheet, min_col=col, min_row=c0, max_row=last), titles_from_data=False)
            s = chart.series[-1]
            s.graphicalProperties.line.solidFill = colour
            s.graphicalProperties.line.width = w
            s.smooth = False
        chart.series[0].tx = SeriesLabel(v=f"{COUNTRIES[c]['name']} sales")
        chart.series[1].tx = SeriesLabel(v="Fitted / forecast")
        chart.set_categories(Reference(fit, min_col=2, min_row=c0, max_row=last))
        fit.add_chart(chart, f"{get_column_letter(fit_width + 2)}{8 + ci * 16}")

    fit.cell(21, 2, "Yellow cells are inputs. Shaded rows at the bottom of the table are the forecast.").font = x.F_MUTED
    x._header(fit, HR, 2, [cfg["unit"].title(), "Status"])
    flu_status = f"'{names['flu']}'!D"
    for i, p in enumerate(periods):
        r = first + i
        fit.cell(r, 2, _period_cells(freq, p)[0]).font = x.F_BASE
        fit.cell(r, 3, f"={flu_status}{r}").font = x.F_BASE
    fit.column_dimensions["B"].width = 46
    fit.column_dimensions["C"].width = 10
    fit.freeze_panes = fit.cell(first, 4)


COUNTRY_CODES = {**{c: c for c in COUNTRIES}, **{v["name"].lower(): c for c, v in COUNTRIES.items()},
                 "norge": "NO", "sverige": "SE", "danmark": "DK"}


def read_sales(path):
    """Read a CSV/XLSX with columns country, week (YYYY-Www) or month (YYYY-MM), sales.

    Returns {(freq, period_key, country_code): value} for prefill()."""
    import pandas as pd

    df = pd.read_csv(path) if str(path).lower().endswith(".csv") else pd.read_excel(path)
    df.columns = [str(c).strip().lower() for c in df.columns]
    if "sales" not in df.columns or "country" not in df.columns or not ({"week", "month"} & set(df.columns)):
        raise ValueError("Sales file needs the columns: country, week or month, sales")
    out = {}
    for _, row in df.iterrows():
        code = COUNTRY_CODES.get(str(row["country"]).strip().lower()) or COUNTRY_CODES.get(str(row["country"]).strip().upper())
        if code is None or pd.isna(row["sales"]):
            continue
        if "week" in df.columns and pd.notna(row.get("week")):
            out[("week", str(row["week"]).strip(), code)] = float(row["sales"])
        if "month" in df.columns and pd.notna(row.get("month")):
            m = row["month"]
            key = f"{m.year}-{m.month:02d}" if hasattr(m, "year") else str(m).strip()[:7]
            out[("month", key, code)] = float(row["sales"])
    return out


def prefill(wb, sales_rows):
    """Write {(freq, period_key, country_code): value} into the input sheets."""
    for freq in FREQ:
        ws = wb[_sheet_names(freq)["input"]]
        rows = {ws.cell(r, 2).value: r for r in range(HR + 1, ws.max_row + 1)}
        for (f, key, code), v in sales_rows.items():
            if f == freq and key in rows:
                ws.cell(rows[key], 4 + list(COUNTRIES).index(code), v)
